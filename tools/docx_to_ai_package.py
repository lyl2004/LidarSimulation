from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from docx import Document

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "o": "urn:schemas-microsoft-com:office:office",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "v": "urn:schemas-microsoft-com:vml",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
}


def qn(prefix: str, local_name: str) -> str:
    return f"{{{NS[prefix]}}}{local_name}"


def sha256_of_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_text(text: str) -> str:
    text = text.replace("\xa0", " ")
    lines = [line.rstrip() for line in text.splitlines()]
    return "\n".join(lines).strip()


def ensure_directory(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def load_xml_from_zip(docx_path: Path, member_name: str) -> ET.Element | None:
    with zipfile.ZipFile(docx_path) as archive:
        try:
            return ET.fromstring(archive.read(member_name))
        except KeyError:
            return None


def load_relationships(docx_path: Path, member_name: str) -> dict[str, dict[str, str]]:
    root = load_xml_from_zip(docx_path, member_name)
    if root is None:
        return {}
    relationships: dict[str, dict[str, str]] = {}
    for rel in root:
        rid = rel.attrib.get("Id", "")
        if rid:
            relationships[rid] = {
                "target": rel.attrib.get("Target", ""),
                "type": rel.attrib.get("Type", ""),
            }
    return relationships


def load_styles(docx_path: Path) -> dict[str, str]:
    root = load_xml_from_zip(docx_path, "word/styles.xml")
    if root is None:
        return {}
    styles: dict[str, str] = {}
    for style in root.findall("w:style", NS):
        style_id = style.attrib.get(qn("w", "styleId"), "")
        name_node = style.find("w:name", NS)
        style_name = ""
        if name_node is not None:
            style_name = name_node.attrib.get(qn("w", "val"), "")
        if style_id:
            styles[style_id] = style_name or style_id
    return styles


def extract_text_with_linebreaks(node: ET.Element) -> str:
    parts: list[str] = []
    for item in node.iter():
        if item.tag in {qn("w", "t"), qn("w", "instrText")}:
            if item.text:
                parts.append(item.text)
        elif item.tag == qn("w", "tab"):
            parts.append("\t")
        elif item.tag in {qn("w", "br"), qn("w", "cr")}:
            parts.append("\n")
    return normalize_text("".join(parts))


def collect_note_refs(node: ET.Element, tag_name: str) -> list[int]:
    values: list[int] = []
    note_tag = qn("w", tag_name)
    attr_name = qn("w", "id")
    for item in node.iter(note_tag):
        raw_value = item.attrib.get(attr_name)
        if raw_value is None:
            continue
        try:
            values.append(int(raw_value))
        except ValueError:
            continue
    return values


def relation_to_artifact(rel_info: dict[str, str], subdir: str) -> str:
    target_name = Path(rel_info.get("target", "")).name
    return f"{subdir}/{target_name}".replace("\\", "/")


def collect_media_refs(node: ET.Element, relationships: dict[str, dict[str, str]]) -> list[dict[str, str]]:
    refs: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()
    rel_attr_embed = qn("r", "embed")
    rel_attr_link = qn("r", "link")
    for item in node.iter():
        rid = ""
        target_subdir = ""
        if item.tag == qn("a", "blip"):
            rid = item.attrib.get(rel_attr_embed, "") or item.attrib.get(rel_attr_link, "")
            target_subdir = "media"
        elif item.tag == qn("v", "imagedata"):
            rid = item.attrib.get(qn("r", "id"), "")
            target_subdir = "media"
        elif item.tag == qn("o", "OLEObject"):
            rid = item.attrib.get(qn("r", "id"), "")
            target_subdir = "embeddings"
        if not rid:
            continue
        rel_info = relationships.get(rid, {})
        target = rel_info.get("target", "")
        if not target:
            continue
        artifact_path = relation_to_artifact(rel_info, target_subdir)
        key = (rid, artifact_path)
        if key in seen:
            continue
        seen.add(key)
        refs.append(
            {
                "relationship_id": rid,
                "target": target,
                "artifact_path": artifact_path,
                "relationship_type": rel_info.get("type", ""),
            }
        )
    return refs


def paragraph_metadata(node: ET.Element, styles: dict[str, str]) -> dict[str, Any]:
    style_id = ""
    style_name = ""
    outline_level: int | None = None
    list_level: int | None = None
    p_style = node.find("w:pPr/w:pStyle", NS)
    if p_style is not None:
        style_id = p_style.attrib.get(qn("w", "val"), "")
        style_name = styles.get(style_id, style_id)
    outline = node.find("w:pPr/w:outlineLvl", NS)
    if outline is not None:
        raw_value = outline.attrib.get(qn("w", "val"))
        if raw_value is not None:
            try:
                outline_level = int(raw_value) + 1
            except ValueError:
                outline_level = None
    num_level = node.find("w:pPr/w:numPr/w:ilvl", NS)
    if num_level is not None:
        raw_value = num_level.attrib.get(qn("w", "val"))
        if raw_value is not None:
            try:
                list_level = int(raw_value)
            except ValueError:
                list_level = None
    return {
        "style_id": style_id,
        "style_name": style_name,
        "outline_level": outline_level,
        "list_level": list_level,
    }


def infer_heading_level(style_id: str, style_name: str, outline_level: int | None) -> int | None:
    if outline_level is not None:
        return max(1, min(outline_level, 6))
    candidates = [style_id, style_name]
    for candidate in candidates:
        match = re.search(r"(heading|标题)\s*([1-6])", candidate, flags=re.IGNORECASE)
        if match:
            return int(match.group(2))
    return None


def build_paragraph_block(
    node: ET.Element,
    styles: dict[str, str],
    relationships: dict[str, dict[str, str]],
) -> dict[str, Any]:
    metadata = paragraph_metadata(node, styles)
    images = []
    objects = []
    for ref in collect_media_refs(node, relationships):
        artifact_path = ref["artifact_path"]
        if artifact_path.startswith("media/"):
            images.append(ref)
        elif artifact_path.startswith("embeddings/"):
            objects.append(ref)
    return {
        "type": "paragraph",
        "text": extract_text_with_linebreaks(node),
        "style_id": metadata["style_id"],
        "style_name": metadata["style_name"],
        "heading_level": infer_heading_level(
            metadata["style_id"],
            metadata["style_name"],
            metadata["outline_level"],
        ),
        "list_level": metadata["list_level"],
        "footnote_ids": collect_note_refs(node, "footnoteReference"),
        "endnote_ids": collect_note_refs(node, "endnoteReference"),
        "images": images,
        "embedded_objects": objects,
    }


def extract_table_cell(node: ET.Element, relationships: dict[str, dict[str, str]]) -> dict[str, Any]:
    colspan = 1
    grid_span = node.find("w:tcPr/w:gridSpan", NS)
    if grid_span is not None:
        raw_value = grid_span.attrib.get(qn("w", "val"))
        if raw_value:
            try:
                colspan = max(1, int(raw_value))
            except ValueError:
                colspan = 1
    v_merge = None
    merge_node = node.find("w:tcPr/w:vMerge", NS)
    if merge_node is not None:
        merge_value = merge_node.attrib.get(qn("w", "val"), "continue")
        v_merge = "restart" if merge_value == "restart" else "continue"
    cell_text = normalize_text("\n\n".join(filter(None, [extract_text_with_linebreaks(par) for par in node.findall("w:p", NS)])))
    images = []
    objects = []
    for ref in collect_media_refs(node, relationships):
        artifact_path = ref["artifact_path"]
        if artifact_path.startswith("media/"):
            images.append(ref)
        elif artifact_path.startswith("embeddings/"):
            objects.append(ref)
    return {
        "text": cell_text,
        "colspan": colspan,
        "v_merge": v_merge,
        "footnote_ids": collect_note_refs(node, "footnoteReference"),
        "endnote_ids": collect_note_refs(node, "endnoteReference"),
        "images": images,
        "embedded_objects": objects,
    }


def compute_table_layout(rows: list[list[dict[str, Any]]]) -> tuple[list[list[dict[str, Any]]], int]:
    active_rowspans: dict[int, dict[str, Any]] = {}
    rendered_rows: list[list[dict[str, Any]]] = []
    max_columns = 0
    for row in rows:
        current_row: list[dict[str, Any]] = []
        next_rowspans: dict[int, dict[str, Any]] = {}
        col_index = 0
        for cell in row:
            if cell["v_merge"] == "continue":
                while col_index not in active_rowspans and active_rowspans:
                    col_index += 1
                origin = active_rowspans.get(col_index)
                if origin is not None:
                    origin["rowspan"] += 1
                    for span_offset in range(origin["colspan"]):
                        next_rowspans[col_index + span_offset] = origin
                col_index += cell["colspan"]
                continue
            while col_index in active_rowspans:
                col_index += 1
            rendered = {
                "text": cell["text"],
                "colspan": cell["colspan"],
                "rowspan": 1,
                "images": cell["images"],
                "embedded_objects": cell["embedded_objects"],
                "footnote_ids": cell["footnote_ids"],
                "endnote_ids": cell["endnote_ids"],
            }
            current_row.append(rendered)
            if cell["v_merge"] == "restart":
                for span_offset in range(cell["colspan"]):
                    next_rowspans[col_index + span_offset] = rendered
            col_index += cell["colspan"]
        if active_rowspans:
            while col_index in active_rowspans:
                col_index += 1
        max_columns = max(max_columns, col_index)
        rendered_rows.append(current_row)
        active_rowspans = next_rowspans
    return rendered_rows, max_columns


def build_table_block(node: ET.Element, relationships: dict[str, dict[str, str]], table_index: int) -> dict[str, Any]:
    raw_rows: list[list[dict[str, Any]]] = []
    for row in node.findall("w:tr", NS):
        raw_rows.append([extract_table_cell(cell, relationships) for cell in row.findall("w:tc", NS)])
    rendered_rows, column_count = compute_table_layout(raw_rows)
    return {
        "type": "table",
        "index": table_index,
        "row_count": len(raw_rows),
        "column_count": column_count,
        "rows": raw_rows,
        "rendered_rows": rendered_rows,
    }


def parse_document_blocks(docx_path: Path) -> list[dict[str, Any]]:
    document_root = load_xml_from_zip(docx_path, "word/document.xml")
    if document_root is None:
        return []
    body = document_root.find("w:body", NS)
    if body is None:
        return []
    relationships = load_relationships(docx_path, "word/_rels/document.xml.rels")
    styles = load_styles(docx_path)
    blocks: list[dict[str, Any]] = []
    table_index = 1
    for child in body:
        if child.tag == qn("w", "p"):
            blocks.append(build_paragraph_block(child, styles, relationships))
        elif child.tag == qn("w", "tbl"):
            blocks.append(build_table_block(child, relationships, table_index))
            table_index += 1
    return blocks


def inspect_docx(path: Path) -> dict[str, Any]:
    with zipfile.ZipFile(path) as archive:
        names = set(archive.namelist())
        bad_member = archive.testzip()
        document_xml_parseable = False
        if "word/document.xml" in names:
            ET.fromstring(archive.read("word/document.xml"))
            document_xml_parseable = True
    doc = Document(path)
    props = doc.core_properties
    return {
        "path": str(path),
        "bytes": path.stat().st_size,
        "sha256": sha256_of_file(path),
        "zip_integrity": bad_member is None,
        "document_xml_parseable": document_xml_parseable,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "core_properties": {
            "title": props.title or "",
            "author": props.author or "",
            "subject": props.subject or "",
            "keywords": props.keywords or "",
        },
        "package_parts": {
            "styles": "word/styles.xml" in names,
            "numbering": "word/numbering.xml" in names,
            "settings": "word/settings.xml" in names,
            "footnotes": "word/footnotes.xml" in names,
            "endnotes": "word/endnotes.xml" in names,
            "comments": "word/comments.xml" in names,
            "media": any(name.startswith("word/media/") for name in names),
            "embeddings": any(name.startswith("word/embeddings/") for name in names),
        },
    }


def collect_part_paragraphs(root: ET.Element) -> list[str]:
    paragraphs: list[str] = []
    for paragraph in root.findall(".//w:p", NS):
        text = extract_text_with_linebreaks(paragraph)
        if text:
            paragraphs.append(text)
    return paragraphs


def extract_headers_and_footers(docx_path: Path) -> dict[str, list[str]]:
    result: dict[str, list[str]] = {}
    with zipfile.ZipFile(docx_path) as archive:
        for member_name in sorted(archive.namelist()):
            if not member_name.startswith("word/"):
                continue
            file_name = Path(member_name).name
            if not (file_name.startswith("header") or file_name.startswith("footer")):
                continue
            if not file_name.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(member_name))
            result[file_name] = collect_part_paragraphs(root)
    return result


def extract_note_part(docx_path: Path, member_name: str, element_name: str) -> list[dict[str, Any]]:
    root = load_xml_from_zip(docx_path, member_name)
    if root is None:
        return []
    items: list[dict[str, Any]] = []
    attr_id = qn("w", "id")
    attr_type = qn("w", "type")
    for note in root.findall(f"w:{element_name}", NS):
        note_type = note.attrib.get(attr_type, "")
        if note_type:
            continue
        raw_id = note.attrib.get(attr_id, "")
        try:
            note_id = int(raw_id)
        except ValueError:
            continue
        paragraphs = collect_part_paragraphs(note)
        items.append(
            {
                "id": note_id,
                "paragraphs": paragraphs,
                "text": "\n\n".join(paragraphs).strip(),
            }
        )
    return items


def extract_binary_parts(docx_path: Path, output_dir: Path, prefix: str, target_dir_name: str) -> list[str]:
    extracted: list[str] = []
    target_dir = output_dir / target_dir_name
    ensure_directory(target_dir)
    with zipfile.ZipFile(docx_path) as archive:
        for member_name in sorted(archive.namelist()):
            if not member_name.startswith(prefix):
                continue
            file_name = Path(member_name).name
            output_path = target_dir / file_name
            with archive.open(member_name) as source_handle, output_path.open("wb") as target_handle:
                target_handle.write(source_handle.read())
            extracted.append(f"{target_dir_name}/{file_name}")
    return extracted


def extract_full_package(docx_path: Path, output_dir: Path) -> None:
    package_dir = output_dir / "package"
    ensure_directory(package_dir)
    with zipfile.ZipFile(docx_path) as archive:
        archive.extractall(package_dir)


def render_paragraph_markdown(block: dict[str, Any]) -> str:
    text = block["text"]
    heading_level = block["heading_level"]
    list_level = block["list_level"]
    lines: list[str] = []
    suffixes: list[str] = []
    if block["footnote_ids"]:
        suffixes.append("footnotes:" + ",".join(str(item) for item in block["footnote_ids"]))
    if block["endnote_ids"]:
        suffixes.append("endnotes:" + ",".join(str(item) for item in block["endnote_ids"]))
    if suffixes:
        text = (text + " " if text else "") + "[" + " | ".join(suffixes) + "]"
    if heading_level is not None and text:
        lines.append(f"{'#' * heading_level} {text}")
    elif list_level is not None and text:
        indent = "  " * max(0, list_level)
        lines.append(f"{indent}- {text}")
    elif text:
        lines.append(text)
    for image in block["images"]:
        lines.append(f"![{Path(image['artifact_path']).name}]({image['artifact_path']})")
    for obj in block["embedded_objects"]:
        lines.append(f"[Embedded object: {obj['artifact_path']}]")
    return "\n".join(lines).strip()


def render_table_cell_text(cell: dict[str, Any]) -> str:
    parts: list[str] = []
    if cell["text"]:
        parts.append(cell["text"])
    for image in cell["images"]:
        parts.append(f"[Image: {image['artifact_path']}]")
    for obj in cell["embedded_objects"]:
        parts.append(f"[Embedded object: {obj['artifact_path']}]")
    if cell["footnote_ids"]:
        parts.append("Footnotes: " + ", ".join(str(item) for item in cell["footnote_ids"]))
    if cell["endnote_ids"]:
        parts.append("Endnotes: " + ", ".join(str(item) for item in cell["endnote_ids"]))
    return "\n".join(parts).strip()


def render_table_markdown(block: dict[str, Any]) -> str:
    lines = [f"<!-- table {block['index']} -->", "<table>"]
    for row in block["rendered_rows"]:
        lines.append("  <tr>")
        for cell in row:
            attrs: list[str] = []
            if cell["colspan"] > 1:
                attrs.append(f' colspan="{cell["colspan"]}"')
            if cell["rowspan"] > 1:
                attrs.append(f' rowspan="{cell["rowspan"]}"')
            cell_text = html.escape(render_table_cell_text(cell)).replace("\n", "<br/>")
            lines.append(f"    <td{''.join(attrs)}>{cell_text}</td>")
        lines.append("  </tr>")
    lines.append("</table>")
    return "\n".join(lines)


def build_markdown_document(source_name: str, blocks: list[dict[str, Any]], notes: dict[str, Any]) -> str:
    parts = [f"# {source_name}"]
    for block in blocks:
        if block["type"] == "paragraph":
            content = render_paragraph_markdown(block)
        else:
            content = render_table_markdown(block)
        if content:
            parts.append(content)
    if notes["headers"]:
        parts.append("## Headers")
        for part_name, paragraphs in notes["headers"].items():
            parts.append(f"### {part_name}")
            parts.extend(paragraphs)
    if notes["footers"]:
        parts.append("## Footers")
        for part_name, paragraphs in notes["footers"].items():
            parts.append(f"### {part_name}")
            parts.extend(paragraphs)
    if notes["footnotes"]:
        parts.append("## Footnotes")
        for item in notes["footnotes"]:
            parts.append(f"### Footnote {item['id']}")
            if item["text"]:
                parts.append(item["text"])
    if notes["endnotes"]:
        parts.append("## Endnotes")
        for item in notes["endnotes"]:
            parts.append(f"### Endnote {item['id']}")
            if item["text"]:
                parts.append(item["text"])
    return "\n\n".join(parts).strip() + "\n"


def build_plain_text(blocks: list[dict[str, Any]], notes: dict[str, Any]) -> str:
    lines: list[str] = []
    for block in blocks:
        if block["type"] == "paragraph":
            text = block["text"]
            if text:
                lines.append(text)
            for image in block["images"]:
                lines.append(f"[Image] {image['artifact_path']}")
            for obj in block["embedded_objects"]:
                lines.append(f"[Embedded object] {obj['artifact_path']}")
        else:
            lines.append(f"[Table {block['index']}]")
            for row in block["rendered_rows"]:
                row_values = [render_table_cell_text(cell).replace("\n", " | ") for cell in row]
                lines.append(" || ".join(row_values))
    if notes["headers"]:
        lines.append("[Headers]")
        for part_name, paragraphs in notes["headers"].items():
            lines.append(part_name)
            lines.extend(paragraphs)
    if notes["footers"]:
        lines.append("[Footers]")
        for part_name, paragraphs in notes["footers"].items():
            lines.append(part_name)
            lines.extend(paragraphs)
    if notes["footnotes"]:
        lines.append("[Footnotes]")
        for item in notes["footnotes"]:
            lines.append(f"{item['id']}: {item['text']}")
    if notes["endnotes"]:
        lines.append("[Endnotes]")
        for item in notes["endnotes"]:
            lines.append(f"{item['id']}: {item['text']}")
    return "\n\n".join(line for line in lines if line).strip() + "\n"


def run_pandoc(source_path: Path, output_dir: Path) -> dict[str, str]:
    result: dict[str, str] = {}
    md_path = output_dir / "pandoc_document.md"
    ast_path = output_dir / "pandoc_ast.json"
    media_dir = output_dir / "pandoc_media"
    commands = [
        (
            "markdown",
            [
                "pandoc",
                str(source_path),
                "-f",
                "docx",
                "-t",
                "gfm+raw_html",
                "--wrap=none",
                "--extract-media",
                str(media_dir),
                "-o",
                str(md_path),
            ],
        ),
        (
            "ast",
            [
                "pandoc",
                str(source_path),
                "-f",
                "docx",
                "-t",
                "json",
                "-o",
                str(ast_path),
            ],
        ),
    ]
    for name, command in commands:
        try:
            subprocess.run(command, check=True, capture_output=True, text=True)
            result[name] = "ok"
        except subprocess.CalledProcessError as exc:
            result[name] = (exc.stderr or exc.stdout or str(exc)).strip()
    return result


def sanitize_name(raw_name: str) -> str:
    cleaned = re.sub(r"[^a-z0-9]+", "_", raw_name.lower()).strip("_")
    return cleaned or "document"


def process_document(source_path: Path, output_dir: Path) -> dict[str, Any]:
    ensure_directory(output_dir)
    inspect_payload = inspect_docx(source_path)
    blocks = parse_document_blocks(source_path)
    notes = {
        "headers": {k: v for k, v in extract_headers_and_footers(source_path).items() if k.startswith("header")},
        "footers": {k: v for k, v in extract_headers_and_footers(source_path).items() if k.startswith("footer")},
        "footnotes": extract_note_part(source_path, "word/footnotes.xml", "footnote"),
        "endnotes": extract_note_part(source_path, "word/endnotes.xml", "endnote"),
    }
    media_files = extract_binary_parts(source_path, output_dir, "word/media/", "media")
    embedding_files = extract_binary_parts(source_path, output_dir, "word/embeddings/", "embeddings")
    extract_full_package(source_path, output_dir)
    pandoc_status = run_pandoc(source_path, output_dir)
    markdown = build_markdown_document(source_path.name, blocks, notes)
    plain_text = build_plain_text(blocks, notes)
    (output_dir / "document.md").write_text(markdown, encoding="utf-8")
    (output_dir / "plain_text.txt").write_text(plain_text, encoding="utf-8")
    write_json(output_dir / "inspect.json", inspect_payload)
    write_json(output_dir / "content.json", {"source_name": source_path.name, "blocks": blocks})
    write_json(output_dir / "notes.json", notes)
    manifest = {
        "source_name": source_path.name,
        "source_path": str(source_path),
        "output_dir": str(output_dir),
        "artifacts": {
            "document_markdown": "document.md",
            "plain_text": "plain_text.txt",
            "content_json": "content.json",
            "notes_json": "notes.json",
            "inspect_json": "inspect.json",
            "package_dir": "package",
            "media_dir": "media",
            "embeddings_dir": "embeddings",
            "pandoc_markdown": "pandoc_document.md",
            "pandoc_ast": "pandoc_ast.json",
            "pandoc_media_dir": "pandoc_media",
        },
        "counts": {
            "blocks": len(blocks),
            "paragraph_blocks": sum(1 for block in blocks if block["type"] == "paragraph"),
            "table_blocks": sum(1 for block in blocks if block["type"] == "table"),
            "media_files": len(media_files),
            "embedding_files": len(embedding_files),
            "footnotes": len(notes["footnotes"]),
            "endnotes": len(notes["endnotes"]),
        },
        "pandoc_status": pandoc_status,
    }
    write_json(output_dir / "manifest.json", manifest)
    return manifest


def parse_entry(raw_entry: str) -> tuple[str, Path]:
    if "=" not in raw_entry:
        raise ValueError(f"Invalid --doc entry: {raw_entry}")
    name, raw_path = raw_entry.split("=", 1)
    clean_name = sanitize_name(name)
    source_path = Path(raw_path).expanduser().resolve()
    return clean_name, source_path


def write_root_readme(output_root: Path, manifests: list[dict[str, Any]]) -> None:
    lines = ["# DOCX AI Extraction", ""]
    for manifest in manifests:
        lines.append(f"## {manifest['source_name']}")
        lines.append(f"- Output directory: `{Path(manifest['output_dir']).name}`")
        lines.append("- Key files: `document.md`, `plain_text.txt`, `content.json`, `notes.json`, `inspect.json`")
        lines.append("- Full OOXML package: `package/`")
        lines.append("")
    (output_root / "README.md").write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract DOCX files into AI-friendly package directories.")
    parser.add_argument("--out-root", required=True, help="Directory for generated output packages")
    parser.add_argument(
        "--doc",
        action="append",
        required=True,
        help="Entry in the form alias=path/to/file.docx",
    )
    args = parser.parse_args()

    output_root = Path(args.out_root).expanduser().resolve()
    ensure_directory(output_root)

    manifests: list[dict[str, Any]] = []
    index_entries: list[dict[str, str]] = []
    for raw_entry in args.doc:
        alias, source_path = parse_entry(raw_entry)
        if not source_path.exists():
            raise FileNotFoundError(source_path)
        if source_path.suffix.lower() != ".docx":
            raise ValueError(f"Expected DOCX file: {source_path}")
        output_dir = output_root / alias
        manifest = process_document(source_path, output_dir)
        manifests.append(manifest)
        index_entries.append(
            {
                "alias": alias,
                "source_name": source_path.name,
                "source_path": str(source_path),
                "output_dir": str(output_dir),
            }
        )

    write_json(output_root / "index.json", {"documents": index_entries})
    write_root_readme(output_root, manifests)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        sys.exit(1)
